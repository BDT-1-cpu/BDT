#!/usr/bin/env python3
"""Refresh the website's figures from the Master Playbook.

    python3 tools/update_site.py "Big Dogs Tour - Master Playbook.xlsx" [--check]

1. Recalculates the Playbook (LibreOffice, full recalculation) so every output sheet is current.
2. Rewrites the data inside the pages from it. Currently: Statistics.html and Trophies.html
   (every table, the dashboard, the trophy-set grid and the Stableford explorer).
3. With --check it changes nothing and lists every figure that would change.

The layout of each table (which sheet, row and columns it comes from) is in tools/sitedata_spec.json,
produced by tools/sitedata_learn.py.
"""
import argparse, json, os, re, shutil, subprocess, sys, tempfile, warnings
import openpyxl
warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from sitedata_learn import norm  # noqa: E402

TROPHY_SET = ['Big Dogs Ryder Cup', 'Green Jacket', '2nd', '3rd', 'Shitzu', 'Poodle', 'Rottweiler', 'Bandit', 'Cheat',
              'Seve & Olly', 'Duncan Goodhew', 'Elton Ravo Prize', 'Great Dane', 'St Bernard', 'Goblat']
CABINET_COL = {'Green Jacket': 'Big Dog', 'Bandit': 'Xolo'}


# ---------------------------------------------------------------- Playbook
def recalculated(playbook):
    """A fully recalculated copy of the Playbook (values only), via LibreOffice."""
    tmp = tempfile.mkdtemp(prefix='bdt_')
    prof = os.path.join(tmp, 'profile', 'user')
    os.makedirs(prof)
    open(os.path.join(prof, 'registrymodifications.xcu'), 'w').write(
        '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry" '
        'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
        '</oor:items>')
    src = os.path.join(tmp, 'playbook.xlsx')
    shutil.copyfile(playbook, src)
    out = os.path.join(tmp, 'out')
    subprocess.run(['soffice', f'-env:UserInstallation=file://{os.path.join(tmp, "profile")}', '--headless', '--calc',
                    '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', out, src],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=1800)
    return os.path.join(out, 'playbook.xlsx')


class Book:
    def __init__(self, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        self.g = {n: [list(r) for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}

    def cell(self, sheet, r, c):
        g = self.g[sheet]
        return g[r][c] if r < len(g) and c < len(g[r]) else None


def tidy(v, fmt=''):
    if isinstance(v, str) and fmt != 'pct':
        return v                     # text exactly as in the Playbook (spacing included)
    v = norm(v)
    if fmt == 'pct' and isinstance(v, (int, float)):
        return f'{round(v * 100)}%'
    if isinstance(v, float):
        v = round(v, 2)
        return int(v) if v == int(v) else v
    if fmt == 'str' and isinstance(v, (int, float)):
        return str(v)
    return v


# ---------------------------------------------------------------- tables
def table(book, sp, header=None):
    rows = []
    r = sp['row']
    n = sp['n'] if sp['extent'] == 'fixed' else 10 ** 6
    key = sp['cols'][0]
    while len(rows) < n:
        if sp['extent'] != 'fixed' and norm(book.cell(sp['sheet'], r, key)) == '':
            break
        if r >= len(book.g[sp['sheet']]):
            break
        rows.append([tidy(book.cell(sp['sheet'], r, c), f) for c, f in zip(sp['cols'], sp['fmt'])])
        r += 1
    return rows


def year_columns(book, sp, block):
    """In by-year tables, take each year's column from the Playbook's own heading row."""
    hr = book.g[sp['sheet']][sp['row'] - 1]
    where = {re.sub(r'\*$', '', str(norm(v))): c for c, v in enumerate(hr) if re.fullmatch(r'(19|20)\d\d\*?', str(norm(v)))}
    H = block['headers']
    isyear = [j > 0 and re.fullmatch(r'(19|20)\d\d\*?', str(h)) is not None for j, h in enumerate(H)]
    if len(H) != len(sp['cols']) and any(isyear):
        # the page already has more year columns than the spec was learned with (added by an earlier run):
        # keep the spec's columns before and after the years, and take every year from the heading row
        first = isyear.index(True)
        after = len(H) - 1 - (len(isyear) - 1 - isyear[::-1].index(True))
        ys = [j for j, v in enumerate(sp['cols']) if j > 0 and j >= first and j < len(sp['cols']) - after]
        yfmt = sp['fmt'][ys[0]] if ys else ''
        sp['cols'] = sp['cols'][:first] + [where.get(re.sub(r'\*$', '', str(h)), 0) for h in H[first:len(H) - after]] + \
            (sp['cols'][len(sp['cols']) - after:] if after else [])
        sp['fmt'] = sp['fmt'][:first] + [yfmt] * (len(H) - after - first) + (sp['fmt'][len(sp['fmt']) - after:] if after else [])
    for j, h in enumerate(block['headers']):
        y = re.sub(r'\*$', '', str(h))
        if j and y in where:
            sp['cols'][j] = where[y]


def year_extend(book, sp, block):
    """Wide by-year tables (Stableford, Placings, Handicaps, RC Points...) gain a column for a new tour."""
    H = block['headers']
    yrs = [i for i, h in enumerate(H) if re.fullmatch(r'(19|20)\d\d\*?', str(h))]
    if len(yrs) < 3 or yrs[-1] != len(H) - 1 - (len(H) - 1 - yrs[-1]):
        return
    last = yrs[-1]
    hr = sp['row'] - 1
    c = sp['cols'][last] + 1
    g = book.g[sp['sheet']]
    while c < len(g[hr]) and re.fullmatch(r'(19|20)\d\d\*?', str(norm(g[hr][c]))):
        has_data = any(norm(book.cell(sp['sheet'], r, c)) != '' for r in range(sp['row'], len(g)))
        if not has_data:
            break
        H.insert(last + 1, str(norm(g[hr][c])))
        sp['cols'].insert(last + 1, c)
        sp['fmt'].insert(last + 1, sp['fmt'][last])
        last += 1
        c += 1
    # ...and for earlier tours the Playbook has started to cover (e.g. 2002-2005 finishing positions)
    first = yrs[0]
    c = sp['cols'][first] - 1
    while c >= 0 and re.fullmatch(r'(19|20)\d\d\*?', str(norm(g[hr][c]))):
        if str(norm(g[hr][c])) in H or not any(norm(book.cell(sp['sheet'], r, c)) not in ('', 'N/A')
                                                for r in range(sp['row'], len(g))):
            break
        H.insert(first, str(norm(g[hr][c])))
        sp['cols'].insert(first, c)
        sp['fmt'].insert(first, sp['fmt'][first])
        c -= 1


def rc_players(book):
    """Ryder Cup player records: each player under their main team, with their combined record."""
    g = book.g['RC Players']
    hdr = next(i for i, r in enumerate(g) if norm(r[0]) == 'Name' and norm(r[1]) == 'Player')
    out = {'Bloodhounds': [], 'Retrievers': []}
    for r in g[hdr + 1:]:
        if norm(r[0]) == '':
            break
        both = norm(r[2]) not in ('', 0) and norm(r[19]) not in ('', 0)
        team = {'BH': 'Bloodhounds', 'RT': 'Retrievers'}.get(norm(r[43]), norm(r[43]))
        if team not in out:
            continue
        out[team].append([norm(r[1]) + ('*' if both else '')] + [tidy(r[c]) for c in range(36, 42)] + [tidy(r[42], 'pct')])
    return out


def big_dogs_ryder_cup(book):
    rc = book.g['RC Results']
    rs = {norm(r[0]): r for r in book.g['Results Summary'] if isinstance(norm(r[0]), int)}
    hdr = next(i for i, r in enumerate(rc) if norm(r[0]) == 'Year' and norm(r[1]) == 'Location')
    rows = []
    for r in rc[hdr + 1:]:
        y = norm(r[0])
        if not isinstance(y, int):
            break
        s = rs.get(y, [None] * 9)
        rows.append([y, norm(r[1]), norm(r[10]), norm(s[7]), norm(r[11]), norm(s[8])])
    return rows


def countback_notes(book, rows):
    """Green Jacket notes: say exactly who tied and on what (the Playbook only says 'decided on a countback')."""
    ind = [r for r in book.g['Individual'] if isinstance(norm(r[0]), int)]
    out = []
    for row in rows:
        y, note = row[0], row[-1]
        if isinstance(note, str) and 'countback' in note.lower():
            top = sorted([(norm(r[7]), norm(r[5])) for r in ind if norm(r[0]) == y and isinstance(norm(r[7]), (int, float))],
                         key=lambda t: -t[0])
            label = {1: 'Big Dog', 2: '2nd', 3: '3rd', 4: '4th'}
            for pts in sorted({t[0] for t in top[:4]}, reverse=True):
                pl = sorted(p for t, p in top[:4] if t == pts and isinstance(p, int))
                if len(pl) > 1:
                    note = ' / '.join(label.get(p, f'{p}th') for p in pl) + f' level on {int(pts)} pts – decided on countback'
                    break
        out.append(row[:-1] + [note])
    return out


# ---------------------------------------------------------------- other data
def dashboard(book):
    g = book.g['Dashboard']
    find = lambda text: next((i, j) for i, r in enumerate(g) for j, v in enumerate(r) if norm(v) == text)
    i, j = find('THE SERIES')
    series = []
    for r in g[i + 1:]:
        if norm(r[j]) in ('', 'Data checks'):
            break
        series.append([norm(r[j]), tidy(r[j + 1])])
    i, j = find('BIG DOG ROLL OF HONOUR (latest first)')
    roll = []
    for r in g[i + 2:]:
        if norm(r[j]) == '' or len(roll) == 12:
            break
        roll.append([norm(r[j]), norm(r[j + 1]), norm(r[j + 2])])
    i, j = find('MOST BIG DOG WINS')
    i2, j2 = find('MOST RYDER CUP POINTS (credited)')
    wins = [[norm(g[k][j]), tidy(g[k][j + 1])] for k in range(i + 1, i + 9) if norm(g[k][j]) != '']
    pts = [[norm(g[k][j2]), tidy(g[k][j2 + 1])] for k in range(i2 + 1, i2 + 9) if norm(g[k][j2]) != '']
    return {'series': series, 'roll_of_honour': roll, 'most_big_dog_wins': wins, 'most_rc_points': pts}


def trophy_closeness(book, old):
    g = book.g['Trophy Cabinet']
    hdr = next(i for i, r in enumerate(g) if norm(r[0]) == 'Name' and norm(r[1]) == 'Player')
    H = [norm(x) for x in g[hdr]]
    # Ryder Cup winners: anyone who played a match for the side that won or kept the cup (see team_data.cup_wins)
    from team_data import cup_wins
    winners = set(cup_wins(book))
    players = []
    for r in g[hdr + 1:]:
        if norm(r[0]) == '':
            break
        have = []
        for t in TROPHY_SET:
            if t == 'Big Dogs Ryder Cup':
                have.append(norm(r[0]) in winners)
            else:
                col = H.index(CABINET_COL.get(t, t))
                have.append(isinstance(norm(r[col]), (int, float)) and norm(r[col]) > 0)
        players.append({'player': norm(r[0]), 'count': sum(have), 'have': have})
    players.sort(key=lambda p: -p['count'])
    # keep the page's tie order where counts are equal (stable sort above keeps sheet order otherwise)
    order = {p['player']: i for i, p in enumerate(old.get('players', []))}
    players.sort(key=lambda p: (-p['count'], order.get(p['player'], 999)))
    return {'order': TROPHY_SET, 'players': players}


def stableford_explorer(book, sheets, old):
    st = sheets['Stableford']['blocks'][0]
    years = [h for h in st['headers'][1:] if re.fullmatch(r'\d{4}', str(h))]
    caps = {r[0]: r[1] for r in sheets['Caps']['blocks'][0]['rows']}
    players = []
    for row in st['rows']:
        vals = dict(zip(st['headers'][1:], row[1:]))
        players.append({'n': row[0], 's': [vals.get(y) if vals.get(y) not in ('', None, 'N/A', '-') else None for y in years],
                        'tours': caps.get(row[0], 0)})
    return {'years': years, 'preYears': old.get('preYears', []), 'players': players}


# ---------------------------------------------------------------- pages
WEBDATA = re.compile(r'(<script id="webdata" type="application/json">)(.*?)(</script>)', re.S)


def rebuild_page(page, book, spec):
    s = open(page, encoding='utf-8').read()
    m = WEBDATA.search(s)
    W = json.loads(m.group(2))
    name = os.path.basename(page)
    for key, sh in W['sheets'].items():
        for bi, block in enumerate(sh['blocks']):
            sp = spec.get(name, {}).get(key, {}).get(str(bi))
            if key == 'RC Players':
                block['rows'] = rc_players(book)[block['label']]
            elif key == 'Big Dogs Ryder Cup':
                block['rows'] = big_dogs_ryder_cup(book)
            elif sp:
                sp = json.loads(json.dumps(sp))
                year_columns(book, sp, block)
                year_extend(book, sp, block)
                rows = table(book, sp)
                if key == 'Green Jacket':
                    rows = countback_notes(book, rows)
                block['rows'] = rows
    if 'dashboard' in W:
        W['dashboard'] = dashboard(book)
    if 'trophyCloseness' in W:
        W['trophyCloseness'] = trophy_closeness(book, W['trophyCloseness'])
    if 'stablefordExplorer' in W:
        W['stablefordExplorer'] = stableford_explorer(book, W['sheets'], W['stablefordExplorer'])
    compact = '", "' not in m.group(2)[:2000]
    text = json.dumps(W, ensure_ascii=False, separators=(',', ':') if compact else None)
    return s[:m.start(2)] + text.replace('</', '<\\/') + s[m.end(2):]


def diff(old, new, path=''):
    out = []
    if isinstance(old, dict) and isinstance(new, dict):
        for k in sorted(set(old) | set(new), key=str):
            if k == 'playerPhotos':
                continue
            out += diff(old.get(k), new.get(k), f'{path}/{k}')
    elif isinstance(old, list) and isinstance(new, list):
        if len(old) != len(new):
            out.append(f'{path}: {len(old)} items -> {len(new)}')
        for i, (a, b) in enumerate(zip(old, new)):
            out += diff(a, b, f'{path}[{i}]')
    elif old != new:
        out.append(f'{path}: {str(old)[:90]!r} -> {str(new)[:90]!r}')
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('playbook')
    ap.add_argument('--check', action='store_true', help='list the changes without writing them')
    ap.add_argument('--recalculated', help='use an already recalculated copy of the Playbook')
    a = ap.parse_args()
    book = Book(a.recalculated or recalculated(a.playbook))
    spec = json.load(open(os.path.join(HERE, 'sitedata_spec.json')))
    for name in ('Statistics.html', 'Trophies.html'):
        page = os.path.join(ROOT, name)
        before = open(page, encoding='utf-8').read()
        after = rebuild_page(page, book, spec)
        old = json.loads(WEBDATA.search(before).group(2))
        new = json.loads(WEBDATA.search(after).group(2))
        changes = diff(old, new)
        print(f'{name}: {len(changes)} change(s)')
        for c in changes:
            print('   ', c)
        if not a.check and changes:
            open(page, 'w', encoding='utf-8').write(after)


if __name__ == '__main__':
    main()
