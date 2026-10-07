"""Build the remaining site pages in the tour-page style (Oct 2026):
Nicknames, Tour Videos, Rules/Privacy/Support and a BDT 2027 'coming soon' page.
Each is a single self-contained HTML file for a Wix Embed HTML box.
Run: python3 newpages/build_newpages.py  (writes newpages/out/*.html)"""
import os, sys, json, html as H
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
import header_a, site_style  # noqa: E402

OUT = os.path.join(HERE, 'out')
os.makedirs(OUT, exist_ok=True)
CREST = open(os.path.join(ROOT, 'home', 'crest.txt')).read().strip()
esc = lambda s: H.escape(str(s), quote=True)

ICONS = {
 'home': '<path d="M3 11l9-7 9 7v9a1 1 0 0 1-1 1h-5v-6h-6v6H4a1 1 0 0 1-1-1z"/>',
 'dog': '<path d="M5 9l-1-5 4 2h8l4-2-1 5"/><path d="M5 9c0 6 3 10 7 10s7-4 7-10"/><circle cx="9.5" cy="12" r=".8"/><circle cx="14.5" cy="12" r=".8"/><path d="M11 15.5h2l-1 1z"/>',
 'calendar': '<rect x="3" y="5" width="18" height="16" rx="2"/><path d="M3 10h18M8 3v4M16 3v4"/>',
 'film': '<rect x="3" y="4" width="18" height="16" rx="2"/><path d="M7 4v16M17 4v16M3 9h4M3 15h4M17 9h4M17 15h4"/>',
 'play': '<rect x="3" y="6" width="13" height="12" rx="2"/><path d="M16 10l5-3v10l-5-3"/>',
 'review': '<path d="M4 5h16v11H8l-4 4z"/><path d="M8 9h8M8 12h5"/>',
 'book': '<path d="M4 4h7a3 3 0 0 1 3 3v13a2 2 0 0 0-2-2H4z"/><path d="M20 4h-4a2 2 0 0 0-2 2"/><path d="M20 4v14h-6"/>',
 'shield': '<path d="M12 3l8 3v6c0 4.5-3.4 8-8 9-4.6-1-8-4.5-8-9V6z"/><path d="M9 12l2 2 4-4"/>',
 'mail': '<rect x="3" y="5" width="18" height="14" rx="2"/><path d="M3 7l9 6 9-6"/>',
 'venue': '<path d="M12 21s-7-6.2-7-11a7 7 0 0 1 14 0c0 4.8-7 11-7 11z"/><circle cx="12" cy="10" r="2.5"/>',
 'captains': '<path d="M4 8l4 4 4-7 4 7 4-4-2 11H6z"/>',
 'kennels': '<path d="M3 20V8l9-5 9 5v12"/><path d="M9 20v-6h6v6"/>',
 'courses': '<path d="M7 21V3l10 4-10 4"/><path d="M4 21h12"/>',
 'teams': '<circle cx="8" cy="8" r="3"/><circle cx="16" cy="8" r="3"/><path d="M2 20c0-3.3 2.7-6 6-6s6 2.7 6 6M10 20c0-3.3 2.7-6 6-6s6 2.7 6 6"/>',
 'preview': '<path d="M3 10v4h4l6 5V5L7 10z"/><path d="M17 8a5 5 0 0 1 0 8"/>',
}
icon = lambda k: f'<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.7" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">{ICONS[k]}</svg>'

BASE_CSS = """
:root{--bg:#050B09;--gold:#F0C566;--gold2:#C99A3D;--ink:#F4F1E4;--soft:#A9BAAE;--dim:#72847A;--gb:rgba(240,197,102,.18);--bh:#ff2a2a;--rt:#5068f2}
*{box-sizing:border-box}html,body{margin:0;max-width:100%;overflow-x:clip}
body{background:radial-gradient(1100px 620px at 14% -8%,rgba(240,197,102,.08),transparent 60%),linear-gradient(160deg,#0a1a13,#050B09 60%);background-color:#050B09;color:var(--ink);font:15px/1.55 Archivo,system-ui,sans-serif;min-height:100vh}
a{color:var(--gold)}img{display:block}
.tabbar{position:relative;z-index:20}
.tabs{display:flex;gap:6px;max-width:1180px;margin:0 auto;padding:12px 18px}
.tabs button{flex:1 1 0;min-width:0;display:flex;flex-direction:column;align-items:center;gap:6px;font:600 11px/1.1 Archivo,system-ui,sans-serif;color:#CBD8CE;background:rgba(255,255,255,.03);border:1px solid var(--gb);border-radius:12px;padding:10px 2px 8px;cursor:pointer;white-space:nowrap;transition:background .15s,color .15s,border-color .15s}
.tabs button span{display:block;max-width:100%;overflow:hidden;text-overflow:clip}
.tabs button svg{width:24px;height:24px;color:var(--gold)}
.tabs button:hover{border-color:rgba(240,197,102,.5);color:#fff}
.tabs button.on{background:linear-gradient(100deg,#F0C566,#C99A3D);color:#14210F;border-color:var(--gold)}.tabs button.on svg{color:#14210F}
@media(min-width:761px){.tabs{justify-content:center}.tabs button{max-width:150px}.tabbar{position:sticky;top:0;background:rgba(5,11,9,.9);backdrop-filter:blur(10px);-webkit-backdrop-filter:blur(10px);border-bottom:1px solid rgba(240,197,102,.12)}}
@media(max-width:760px){.tabs{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));padding:12px 10px 14px}.tabs button{font-size:10.5px;padding:10px 2px 8px;white-space:normal;text-align:center}.tabs button svg{width:22px;height:22px}}
main{max-width:1180px;margin:0 auto;padding:18px 18px 60px;min-height:min(100vh,1000px)}
@media(max-width:760px){main{padding:14px 12px 50px}}
.view{display:none}.view.on{display:block}
.sh{display:flex;align-items:baseline;justify-content:space-between;gap:12px;flex-wrap:wrap;margin:6px 0 14px}
.sh h2{margin:0;font:400 34px/1 Anton,Impact,sans-serif;letter-spacing:.02em;text-transform:uppercase;color:#FBF7E9}.sh h2 em{font-style:normal;color:var(--gold)}
.sh .sub{color:var(--soft);font-size:13.5px}
@media(max-width:760px){.sh h2{font-size:26px}}
.card{background:linear-gradient(160deg,#10261c,#0b1913);border:1px solid rgba(240,197,102,.16);border-radius:16px;padding:16px 18px;box-shadow:0 14px 30px -20px rgba(0,0,0,.9)}
.muted{color:var(--soft)}
.chips{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 16px}
.chips button{font:600 12.5px/1.15 Archivo,system-ui,sans-serif;border-radius:11px;border:1px solid rgba(240,197,102,.22);background:#0f2219;color:#E6ECE7;padding:8px 12px;cursor:pointer}
.chips button.on{background:linear-gradient(100deg,#F0C566,#C99A3D);color:#14210F;border-color:#F0C566}
.search{width:100%;max-width:420px;font:15px Archivo,system-ui,sans-serif;color:var(--ink);background:#0b1913;border:1px solid rgba(240,197,102,.25);border-radius:12px;padding:10px 14px;margin:0 0 16px;outline:none}
.search:focus{border-color:var(--gold)}
#toTop{position:fixed;right:14px;bottom:14px;z-index:30;font:700 13px Archivo,system-ui,sans-serif;background:var(--gold);color:#14210F;border:0;border-radius:999px;padding:9px 14px;box-shadow:0 4px 14px rgba(0,0,0,.35);cursor:pointer;opacity:0;pointer-events:none;transition:opacity .2s}
#toTop.on{opacity:.95;pointer-events:auto}
html:not(.np-ready) #tabs,html:not(.np-ready) main{visibility:hidden}
"""

BASE_JS = """
(function(){const $=s=>document.querySelector(s);
const tabs=[...document.querySelectorAll('#tabs button')],views=[...document.querySelectorAll('.view')];
let tabsInView=true;try{new IntersectionObserver(e=>{tabsInView=e[0].isIntersecting},{threshold:0}).observe(document.querySelector('.tabbar'))}catch(e){}
function toTabs(force){const a=$('#tabanchor');const y0=a.getBoundingClientRect().top+scrollY;const mob=matchMedia('(max-width:760px)').matches;
 if(!force&&!mob&&tabsInView&&scrollY<=y0)return;window.scrollTo({top:y0,behavior:'instant'});try{a.scrollIntoView({block:'start',behavior:'instant'})}catch(e){}}
window.goTab=function(k,noScroll){tabs.forEach(b=>b.classList.toggle('on',b.dataset.k===k));views.forEach(v=>v.classList.toggle('on',v.id==='v-'+k));if(!noScroll)toTabs(false);
 try{history.replaceState(null,'','#'+k)}catch(e){}document.dispatchEvent(new CustomEvent('tabshown',{detail:k}))};
tabs.forEach(b=>b.addEventListener('click',()=>goTab(b.dataset.k)));
const h=(location.hash||'').slice(1);if(h&&tabs.some(b=>b.dataset.k===h))goTab(h,true);
function fitTabs(){const t=$('#tabs');tabs.forEach(b=>b.style.fontSize='');if(matchMedia('(max-width:760px)').matches)return;let fs=11;
 const over=()=>tabs.some(b=>{const sp=b.querySelector('span');return sp&&sp.scrollWidth>b.clientWidth-4});while(over()&&fs>8.5){fs-=.25;tabs.forEach(b=>b.style.fontSize=fs+'px')}}
fitTabs();addEventListener('resize',fitTabs);
const tt=$('#toTop');addEventListener('scroll',()=>tt.classList.toggle('on',scrollY>700),{passive:true});tt.onclick=()=>toTabs(true);
document.documentElement.classList.add('np-ready')})();
"""

THEME_CSS = """
.bdt-hero.o-th{background-image:none}
.bdt-hero .thbg{position:absolute;inset:0;z-index:0;overflow:hidden;pointer-events:none}
.bdt-hero.o-th::before{z-index:1;background:linear-gradient(90deg,rgba(5,11,9,.94) 0%,rgba(5,11,9,.66) 45%,rgba(5,11,9,.25) 85%),linear-gradient(0deg,rgba(5,11,9,.85),transparent 55%)}
.bdt-hero.o-th .bdt-hin,.bdt-hero.o-th .bdt-stripe{z-index:2}
.bdt-hero.th-np{background:radial-gradient(ellipse at 80% 30%,#163626,#050B09 70%)}
.bdt-hero .words{position:absolute;inset:-6% -6% -6% 40%;display:flex;flex-wrap:wrap;align-content:center;gap:2px 22px;transform:rotate(-6deg)}
.bdt-hero .words b{font:400 40px/1.05 Anton,Impact,sans-serif;color:transparent;-webkit-text-stroke:1.2px rgba(240,197,102,.3);text-transform:uppercase;white-space:nowrap}
.bdt-hero .mos{position:absolute;inset:0 0 0 30%;display:grid;grid-template-columns:repeat(6,1fr);grid-template-rows:repeat(3,1fr);gap:3px}
.bdt-hero .mos span{overflow:hidden;background:#0d2419}.bdt-hero .mos img{width:100%;height:100%;object-fit:cover;filter:saturate(.9) brightness(.8)}
.bdt-hero .big{position:absolute;right:-2%;top:50%;transform:translateY(-50%) rotate(-6deg);font:400 260px/1 Anton,Impact,sans-serif;color:transparent;-webkit-text-stroke:2px rgba(240,197,102,.28)}
@media(max-width:760px){.bdt-hero.o-th::before{background:linear-gradient(0deg,rgba(5,11,9,.95),rgba(5,11,9,.6) 60%,rgba(5,11,9,.35))}
 .bdt-hero .words{inset:-10% -40% -10% 20%}.bdt-hero .words b{font-size:26px}.bdt-hero .mos{inset:0;grid-template-columns:repeat(4,1fr)}.bdt-hero .big{font-size:150px;right:-8%}}
"""


def page(title, sub, theme_bg, tabs, views, extra_css='', extra_js='', note='', kicker=None, wix_title=''):
    crest = f'<img class="bdt-crest" src="{CREST}" alt="Big Dogs Tour crest">'
    hdr = header_a.header_html(crest, title, sub, '', 'o-th th-np', kicker).replace(' style="--hero:url(\'\')"', '')
    hdr = hdr.replace('<div class="bdt-hin">', f'<div class="thbg">{theme_bg}</div><div class="bdt-hin">', 1)
    if note:
        hdr = hdr.replace('</div></div></div><div class="bdt-stripe">', f'</div><div class="bdt-note">{note}</div></div></div><div class="bdt-stripe">', 1)
    tab_html = ''.join(f'<button data-k="{k}" class="{"on" if i == 0 else ""}">{icon(ic)}<span>{esc(lbl)}</span></button>' for i, (k, lbl, ic) in enumerate(tabs))
    view_html = ''.join(f'<section class="view{" on" if i == 0 else ""}" id="v-{k}">{views[k]}</section>' for i, (k, _, __) in enumerate(tabs))
    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{esc(wix_title or sub)} | Big Dogs Tour</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link href="https://fonts.googleapis.com/css2?family=Anton&family=Archivo:wght@400;600;700&display=swap" rel="stylesheet">
<style>{BASE_CSS}{THEME_CSS}{extra_css}</style>
<script>setTimeout(function(){{document.documentElement.classList.add('np-ready')}},2500)</script>
</head><body>
{hdr}
<div id="tabanchor"></div><nav class="tabbar"><div class="tabs" id="tabs">{tab_html}</div></nav>
<main id="views">{view_html}</main>
<button id="toTop">▲ Tabs</button>
<script>{extra_js}</script>
<script>{BASE_JS}</script>
</body></html>"""
    i = doc.find('</head>')
    return doc[:i] + header_a.CSS + doc[i:]


# ---------------------------------------------------------------- Nicknames
def nicknames():
    rows = None
    src = os.environ.get('BDT_SRC') or os.path.dirname(ROOT)
    pb = os.path.join(src, 'Big Dogs Tour - Master Playbook.xlsx')
    if os.path.exists(pb):  # the playbook's Nicknames sheet is the master copy (added Oct 2026)
        import openpyxl
        wb = openpyxl.load_workbook(pb, read_only=True, data_only=True)
        if 'Nicknames' in wb.sheetnames:
            raw = [[('' if v is None else str(v).strip()) for v in r] for r in wb['Nicknames'].iter_rows(values_only=True)]
            hdr = raw[0]
            ycols = [i for i, h in enumerate(hdr) if h.isdigit()]
            rows = [['DOG'] + [hdr[i] for i in ycols]] + [[r[0]] + [(r[i] if i < len(r) and r[i] else '-') for i in ycols] for r in raw[1:] if r and r[0]]
    if rows is None:
        rows = [l.split('|') for l in open(os.path.join(HERE, 'nicknames.txt'), encoding='utf-8').read().strip().split('\n')]
    years = rows[0][1:]
    dogs = []
    for r in rows[1:]:
        names = [(y, n.strip()) for y, n in zip(years, r[1:]) if n.strip() not in ('', '-')]
        dogs.append({'d': r[0], 'n': names})
    words = [n for d in dogs for _, n in d['n']][::3][:26]
    bg = '<div class="words">' + ''.join(f'<b>{esc(w)}</b>' for w in words) + '</div>'
    total = sum(len(d['n']) for d in dogs)
    data = json.dumps({'years': years, 'dogs': dogs}, ensure_ascii=False)
    views = {
        'dogs': f'<div class="sh"><h2>Dog <em>by</em> Dog</h2><span class="sub">{total} nicknames across {len(dogs)} dogs · type to search</span></div>'
                '<input class="search" id="nkQ" type="search" placeholder="Search a dog or a nickname…" autocomplete="off"><div class="nkgrid" id="nkGrid"></div>',
        'tours': '<div class="sh"><h2>Tour <em>by</em> Tour</h2><span class="sub" id="nkYsub"></span></div><div class="chips" id="nkYears"></div><div class="nklist" id="nkYear"></div>',
    }
    css = """
.nkgrid{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:12px}
.nk{padding:14px 16px}.nk h3{margin:0 0 8px;font:400 22px/1 Anton,Impact,sans-serif;letter-spacing:.02em;color:var(--gold);display:flex;justify-content:space-between;align-items:baseline}
.nk h3 small{font:600 11px Archivo,system-ui,sans-serif;color:var(--soft);letter-spacing:.06em;text-transform:uppercase}
.nk ul{list-style:none;margin:0;padding:0}.nk li{display:flex;gap:10px;padding:4px 0;border-top:1px solid rgba(240,197,102,.08);font-size:14px}
.nk li:first-child{border-top:0}.nk li i{font-style:normal;font:600 12px/1.6 Archivo,system-ui,sans-serif;color:var(--dim);min-width:38px}
.nk.none{opacity:.55}.nk mark{background:rgba(240,197,102,.3);color:inherit;border-radius:3px}
.nklist{display:grid;grid-template-columns:repeat(auto-fill,minmax(230px,1fr));gap:10px}
.nkrow{display:flex;flex-direction:column;gap:2px;padding:12px 14px}.nkrow b{font:400 19px/1.05 Anton,Impact,sans-serif;letter-spacing:.02em;color:#FBF7E9}.nkrow span{font:600 11.5px Archivo,system-ui,sans-serif;color:var(--gold);letter-spacing:.06em;text-transform:uppercase}
.empty{padding:26px;text-align:center;color:var(--soft)}"""
    js = f"""const NK={data};
document.addEventListener('DOMContentLoaded',()=>{{const e=s=>String(s).replace(/[&<>"]/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}}[c]));
const g=document.getElementById('nkGrid'),q=document.getElementById('nkQ');
const hl=(t,s)=>{{if(!s)return e(t);const i=t.toLowerCase().indexOf(s);return i<0?e(t):e(t.slice(0,i))+'<mark>'+e(t.slice(i,i+s.length))+'</mark>'+e(t.slice(i+s.length))}};
function draw(){{const s=(q.value||'').trim().toLowerCase();const L=NK.dogs.filter(d=>!s||d.d.toLowerCase().includes(s)||d.n.some(x=>x[1].toLowerCase().includes(s)));
 g.innerHTML=L.length?L.map(d=>`<div class="card nk${{d.n.length?'':' none'}}"><h3>${{hl(d.d,s)}}<small>${{d.n.length?d.n.length+' nickname'+(d.n.length>1?'s':''):'none on record'}}</small></h3>${{d.n.length?'<ul>'+d.n.map(x=>`<li><i>${{x[0]}}</i><span>${{hl(x[1],s)}}</span></li>`).join('')+'</ul>':''}}</div>`).join(''):'<div class="card empty">No dog or nickname matches that search.</div>'}}
q.addEventListener('input',draw);draw();
const yc=document.getElementById('nkYears'),yl=document.getElementById('nkYear'),ys=document.getElementById('nkYsub');
const has=y=>NK.dogs.filter(d=>d.n.some(x=>x[0]===y));const YS=NK.years.filter(y=>has(y).length);
function showY(y){{yc.querySelectorAll('button').forEach(b=>b.classList.toggle('on',b.dataset.y===y));const L=has(y);ys.textContent=L.length+' dogs with a nickname in '+y;
 yl.innerHTML=L.map(d=>`<div class="card nkrow"><span>${{e(d.d)}}</span><b>${{e(d.n.find(x=>x[0]===y)[1])}}</b></div>`).join('')}}
yc.innerHTML=YS.map(y=>`<button data-y="${{y}}">${{y}}</button>`).join('');yc.addEventListener('click',ev=>{{const b=ev.target.closest('button');if(b)showY(b.dataset.y)}});
showY(YS[YS.length-1]);}});"""
    tabs = [('dogs', 'By Dog', 'dog'), ('tours', 'By Tour', 'calendar')]
    return page('Big Dogs Tour', 'Nicknames', bg, tabs, views, css, js, note='Every tour nickname, dog by dog and tour by tour.', wix_title='Nicknames')


# ---------------------------------------------------------------- Tour videos
def videos():
    V = json.load(open(os.path.join(HERE, 'videos.json'), encoding='utf-8'))
    th = lambda v, w=640: f"https://i.vimeocdn.com/video/{v['th'].rsplit('-d_', 1)[0]}-d_{w}"
    bg = '<div class="mos">' + ''.join(f'<span><img src="{th(v, 295)}" alt="" loading="lazy"></span>' for v in V[:18]) + '</div>'
    import base64
    cdir = os.path.join(HERE, 'covers')
    covers = {}
    for v in V:
        c = v.get('cv')
        if c and c not in covers and os.path.exists(os.path.join(cdir, c + '_front.jpg')):
            enc = lambda n: 'data:image/jpeg;base64,' + base64.b64encode(open(os.path.join(cdir, n), 'rb').read()).decode()
            covers[c] = {'f': enc(c + '_front.jpg'), 'b': enc(c + '_back.jpg') if os.path.exists(os.path.join(cdir, c + '_back.jpg')) else ''}
    data = json.dumps([{**v, 'th': th(v)} for v in V], ensure_ascii=False)
    cov_js = json.dumps(covers)
    yrs = sorted({v['y'] for v in V})
    cnt = lambda f: sum(1 for v in V if f(v))
    views = {
        'all': f'<div class="sh"><h2>Video<em>dog</em>raphy</h2><span class="sub">{len(V)} films from {len(yrs)} tours · tap any film to play</span></div><div class="vidwrap" data-f="all"></div>',
        'previews': f'<div class="sh"><h2>The <em>Previews</em></h2><span class="sub">{cnt(lambda v: v["k"] == "preview")} preview DVDs</span></div><div class="vidwrap" data-f="previews"></div>',
        'reviews': f'<div class="sh"><h2>The <em>Reviews</em></h2><span class="sub">{cnt(lambda v: v["k"] == "review")} review films</span></div><div class="vidwrap" data-f="reviews"></div>',
    }
    css = """
.vy{margin:4px 0 22px}.vy h3{margin:0 0 10px;font:400 22px/1 Anton,Impact,sans-serif;letter-spacing:.04em;color:var(--gold)}
.vgrid{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:14px}
@media(max-width:1000px){.vgrid{grid-template-columns:repeat(3,minmax(0,1fr))}}@media(max-width:760px){.vgrid{grid-template-columns:repeat(2,minmax(0,1fr));gap:10px}.vc .ti{font-size:12.5px;padding:9px 10px 10px}}@media(max-width:330px){.vgrid{grid-template-columns:minmax(0,1fr)}}
.vc .yr{position:absolute;left:8px;top:8px;font:400 14px Anton,Impact,sans-serif;letter-spacing:.06em;color:#14210F;background:var(--gold);padding:2px 8px;border-radius:6px}
.vc{display:flex;flex-direction:column}.vc .ti{flex:1}
.vc{position:relative;padding:0;overflow:hidden;cursor:pointer;text-align:left;font:inherit;color:inherit}
.vc .th{position:relative;aspect-ratio:5/7;background:#0d2419;overflow:hidden}.vc .th img{position:absolute;inset:0;width:100%;height:100%;object-fit:cover}
.vc .th img.bk{opacity:0;transition:opacity .35s}.vc.back .th img.bk{opacity:1}@media(hover:hover){.vc.cov:hover .th img.bk{opacity:1}}
.vc .flip{position:absolute;left:8px;bottom:8px;z-index:2;font:700 11px Archivo,system-ui,sans-serif;background:rgba(0,0,0,.7);color:#fff;padding:3px 8px;border-radius:6px;cursor:pointer}
.vc .pl,.vc .du,.vc .yr{z-index:2}.vc .th:after{content:'';position:absolute;inset:0;background:linear-gradient(180deg,rgba(0,0,0,.15),transparent 30%,transparent 70%,rgba(0,0,0,.45));pointer-events:none;z-index:1}
.vc .pl{position:absolute;left:50%;top:50%;width:54px;height:54px;margin:-27px 0 0 -27px;border-radius:50%;background:rgba(5,11,9,.65);border:2px solid var(--gold);display:grid;place-items:center}
.vc .pl:after{content:"";border-left:16px solid var(--gold);border-top:10px solid transparent;border-bottom:10px solid transparent;margin-left:4px}
.vc .du{position:absolute;right:8px;bottom:8px;font:700 11px Archivo,system-ui,sans-serif;background:rgba(0,0,0,.7);padding:3px 7px;border-radius:6px}
.vc .ti{padding:11px 14px 13px;font-weight:600;font-size:14px}.vc:hover{border-color:rgba(240,197,102,.5)}
.vlb{position:fixed;inset:0;z-index:1000;background:rgba(0,0,0,.92);display:none;align-items:center;justify-content:center;flex-direction:column;padding:18px}
.vlb.on{display:flex}.vlb .fr{width:min(1100px,100%);aspect-ratio:16/9;background:#000}.vlb iframe{width:100%;height:100%;border:0}
.vlb .cap{margin-top:12px;font:400 18px Anton,Impact,sans-serif;letter-spacing:.06em;color:var(--gold);text-align:center}
.vlb .x{position:absolute;top:12px;right:12px;width:44px;height:44px;border-radius:50%;border:0;background:rgba(255,255,255,.14);color:#fff;font:26px/1 Archivo;cursor:pointer}"""
    js = f"""const VID={data};const COV={cov_js};
document.addEventListener('DOMContentLoaded',()=>{{const e=s=>String(s).replace(/[&<>"]/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}}[c]));
const dur=s=>{{const h=Math.floor(s/3600),m=Math.floor(s%3600/60),x=s%60;return (h?h+':'+String(m).padStart(2,'0'):m)+':'+String(x).padStart(2,'0')}};
const F={{all:()=>true,previews:v=>v.k==='preview',reviews:v=>v.k==='review'}};
document.querySelectorAll('.vidwrap').forEach(w=>{{const L=VID.filter(F[w.dataset.f]);const ys=[...new Set(L.map(v=>v.y))].sort((a,b)=>b-a);
 const S=L.slice().sort((a,b)=>b.y-a.y||a.t.localeCompare(b.t,undefined,{{numeric:true}}));w.innerHTML='<div class="vgrid">'+S.map(v=>{{const C=v.cv&&COV[v.cv];return `<button class="card vc${{C?' cov':''}}" data-id="${{v.id}}"><div class="th"><img src="${{C?C.f:v.th}}" alt="" loading="lazy">${{C&&C.b?`<img class="bk" src="${{C.b}}" alt="" loading="lazy"><span class="flip" title="See the back cover">↻ Back</span>`:''}}<span class="pl"></span><span class="yr">${{v.y}}</span><span class="du">${{dur(v.d)}}</span></div><div class="ti">${{e(v.t)}}</div></button>`}}).join('')+'</div>'}});
const lb=document.createElement('div');lb.className='vlb';lb.innerHTML='<button class="x" aria-label="Close">×</button><div class="fr"></div><div class="cap"></div>';document.body.appendChild(lb);
const close=()=>{{lb.classList.remove('on');lb.querySelector('.fr').innerHTML=''}};
document.addEventListener('click',ev=>{{const fl=ev.target.closest('.flip');if(fl){{ev.stopPropagation();const card=fl.closest('.vc');card.classList.toggle('back');fl.textContent=card.classList.contains('back')?'↻ Front':'↻ Back';return}}const c=ev.target.closest('.vc');if(c){{const v=VID.find(x=>x.id===c.dataset.id);lb.querySelector('.fr').innerHTML=`<iframe src="https://player.vimeo.com/video/${{v.id}}?autoplay=1&title=0&byline=0&portrait=0" allow="autoplay; fullscreen; picture-in-picture" allowfullscreen></iframe>`;lb.querySelector('.cap').textContent=v.t;lb.classList.add('on');return}}
 if(ev.target===lb||ev.target.classList.contains('x'))close()}});document.addEventListener('keydown',ev=>{{if(ev.key==='Escape')close()}});}});"""
    tabs = [('all', 'All Films', 'film'), ('previews', 'Previews', 'play'), ('reviews', 'Reviews', 'review')]
    return page('Big Dogs Tour', 'Videos', bg, tabs, views, css, js, note='Preview and review films from the tour archive.', wix_title='Videos')


# ---------------------------------------------------------------- Rules, privacy & support
def rules():
    U = 'https://www.bigdogstour.com/_files/ugd/'
    img = lambda i, w=500, h=700: f'https://static.wixstatic.com/media/588d17_{i}~mv2.jpg/v1/fit/w_{w},h_{h},q_85/i.jpg'
    docs = [('The BDT Rules of Golf', 'The Rules of Golf as varied and supplemented by the BDT Rules of Golf.', U + '588d17_5fc679a96feb48e2bd16ddcae554fb0e.pdf', img('48a4b246a7c34df289f53bd174d42fc0'), 'Open PDF'),
            ('The R&amp;A Rules of Golf', 'The official R&amp;A / USGA Rules of Golf (effective January 2019).', U + '588d17_687a70353a954ba39bb0be6ca12d0150.pdf', img('17533d02c1cb43f8b706e38d56d2ccc5'), 'Open PDF')]
    cards = ''.join(f'<a class="card doc" href="{u}" target="_blank" rel="noopener"><img src="{im}" alt="" loading="lazy"><div><h3>{t}</h3><p>{d}</p><span class="btn">{b} →</span></div></a>' for t, d, u, im, b in docs)
    rules_v = ('<div class="sh"><h2>The BDT <em>Rules</em></h2><span class="sub">How we play – read them before you tee off</span></div>'
               f'<div class="docs">{cards}</div>'
               f'<div class="card dl"><div><h3>BDT Rules (Word version)</h3><p class="muted">Download the editable copy of the BDT rules.</p></div><a class="btn" href="https://github.com/BDT-1-cpu/BDT/raw/main/BDT%20Rules%202021%20Version.doc" target="_blank" rel="noopener">Download →</a></div>')
    body, ul = [], False
    for line in open(os.path.join(HERE, 'privacy.txt'), encoding='utf-8').read().strip().split('\n'):
        if line.startswith('- '):
            if not ul:
                body.append('<ul>'); ul = True
            body.append(f'<li>{line[2:]}</li>'); continue
        if ul:
            body.append('</ul>'); ul = False
        if line.startswith('### '):
            body.append(f'<h4>{line[4:]}</h4>')
        elif line.startswith('## '):
            body.append(f'<h3>{line[3:]}</h3>')
        else:
            body.append(f'<p>{line}</p>')
    if ul:
        body.append('</ul>')
    priv_v = ('<div class="sh"><h2>Website <em>Privacy</em> Policy</h2><span class="sub">How bigdogstour.com and the BDT app handle your data</span></div>'
              f'<div class="card legal">{"".join(body)}</div>')
    contact_v = ('<div class="sh"><h2>Contact <em>Us</em></h2><span class="sub">Questions, corrections or a photo we have missed</span></div>'
                 '<div class="card contact"><p>Spotted a wrong score, a missing photo or a nickname we have forgotten? Use the contact form '
                 '<b>below</b> and the Big Dogs will get back to you.</p></div>')
    views = {'rules': rules_v, 'privacy': priv_v, 'contact': contact_v}
    css = """
.docs{display:grid;grid-template-columns:repeat(auto-fill,minmax(320px,1fr));gap:14px;margin-bottom:14px}
.doc{display:flex;gap:16px;align-items:center;text-decoration:none;color:inherit}.doc img{width:110px;height:auto;border-radius:8px;box-shadow:0 8px 20px rgba(0,0,0,.6);flex:none;background:#fff}
.doc h3,.dl h3{margin:0 0 6px;font:400 22px/1.05 Anton,Impact,sans-serif;letter-spacing:.02em;color:#FBF7E9}.doc p{margin:0 0 10px;color:var(--soft);font-size:14px}
.btn{display:inline-block;font:700 13px Archivo,system-ui,sans-serif;color:#14210F;background:linear-gradient(100deg,#F0C566,#C99A3D);border-radius:10px;padding:8px 14px;text-decoration:none;white-space:nowrap}
.doc:hover{border-color:rgba(240,197,102,.5)}
.dl{display:flex;align-items:center;justify-content:space-between;gap:14px;flex-wrap:wrap}.dl p{margin:0}
.legal{padding:22px 26px;max-width:900px}.legal h3{font:400 22px/1.1 Anton,Impact,sans-serif;letter-spacing:.03em;color:var(--gold);margin:22px 0 8px}.legal h3:first-child{margin-top:0}
.legal h4{font:700 15px Archivo,system-ui,sans-serif;color:#FBF7E9;margin:16px 0 6px}.legal p,.legal li{color:#D9E2DB;font-size:14.5px;line-height:1.6}.legal p{margin:0 0 10px}.legal ul{margin:0 0 12px;padding-left:20px}
.contact{max-width:760px}.contact p{margin:0;font-size:15px}
@media(max-width:760px){.legal{padding:16px}.doc img{width:84px}}"""
    tabs = [('rules', 'BDT Rules', 'book'), ('privacy', 'Privacy', 'shield'), ('contact', 'Contact', 'mail')]
    bg = '<div class="words">' + ''.join(f'<b>{w}</b>' for w in ['Rule 1', 'Rule 2', 'Play it as it lies', 'Rule 13', 'Gimme', 'Rule 18', 'Mulligan?', 'Rule 9', 'Free drop', 'Rule 14', 'Penalty', 'Rule 17', 'Out of bounds', 'Rule 6']) + '</div>'
    return page('Big Dogs Tour', 'Rules, Privacy &amp; Support', bg, tabs, views, css, '', note='The BDT rules of golf, our privacy policy and how to get in touch.', wix_title='Rules, Privacy & Support')


# ---------------------------------------------------------------- BDT 2027 (coming soon)
def tour_2027():
    soon = lambda h, t: f'<div class="sh"><h2>{h}</h2></div><div class="card soon"><b>Coming soon</b><p>{t}</p></div>'
    views = {
        'overview': ('<div class="card soon big"><span class="pill">The next Big Dogs Tour</span><b>2027</b>'
                     '<p>The planning is under way. Destination, dates, captains and kennels will appear here as soon as they are confirmed.</p></div>'),
        'venue': soon('The <em>Venue</em>', 'The destination will be revealed here.'),
        'captains': soon('The <em>Captains</em>', 'The 2027 Bloodhound and Retriever captains will be named here.'),
        'kennels': soon('The <em>Accommodation</em>', 'Photos of the 2027 kennels will appear here once booked.'),
        'courses': soon('The Golf <em>Courses</em>', 'The courses for all three rounds will be listed here.'),
        'itinerary': soon('The <em>Itinerary</em>', 'Day-by-day plans will appear here nearer the time.'),
        'teams': soon('The <em>Teams</em>', 'The Bloodhounds and Retrievers line-ups will appear after the draft.'),
        'announcement': soon('The Team <em>Announcement</em>', 'The team announcement video will go here.'),
    }
    css = """.soon{padding:26px 24px}.soon b{display:block;font:400 26px/1 Anton,Impact,sans-serif;letter-spacing:.04em;color:var(--gold);text-transform:uppercase;margin-bottom:8px}.soon p{margin:0;color:var(--soft);max-width:640px}
.soon.big{text-align:center;padding:46px 24px}.soon.big b{font-size:110px;color:#FBF7E9;margin:12px 0}.soon.big p{margin:0 auto}
.pill{display:inline-block;font:700 11px Archivo,system-ui,sans-serif;letter-spacing:.14em;text-transform:uppercase;color:var(--gold);border:1px solid rgba(240,197,102,.4);border-radius:999px;padding:5px 12px}
@media(max-width:760px){.soon.big b{font-size:76px}}"""
    tabs = [('overview', 'Overview', 'home'), ('venue', 'Venue', 'venue'), ('captains', 'Captains', 'captains'), ('kennels', 'Kennels', 'kennels'),
            ('courses', 'Courses', 'courses'), ('itinerary', 'Itinerary', 'calendar'), ('teams', 'Teams', 'teams'), ('announcement', 'Preview', 'preview')]
    bg = '<div class="big">2027</div>'
    hd = page('Big Dogs Tour <span class="bdt-yr">2027</span>', 'Destination to be announced', bg, tabs, views, css, '',
              note='', wix_title='BDT 2027')
    cal = ('<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round">'
           '<rect x="3" y="5" width="18" height="16" rx="2"/><path d="M3 10h18M8 3v4M16 3v4"/></svg>')
    extra = (f'<div class="bdt-strap">Coming soon</div><div class="bdt-when">{cal}<span>Dates to be confirmed</span></div>')
    hd = hd.replace('<div class="bdt-t2">Destination to be announced</div>', '<div class="bdt-t2">Destination to be announced</div>' + extra, 1)
    strap_css = """<style>.bdt-hero .bdt-strap{margin-top:9px;padding-top:8px;border-top:1px solid rgba(240,197,102,.25);font:17px/1 Anton,Impact,sans-serif;letter-spacing:.05em;color:#FBF7E9;text-transform:uppercase}
.bdt-hero .bdt-when{display:flex;align-items:center;gap:6px;margin-top:6px;font:600 11.5px/1.25 Archivo,system-ui,sans-serif;color:#d6dfd8}.bdt-hero .bdt-when svg{width:13px;height:13px;color:#F0C566}
@media(min-width:761px){.bdt-hero .bdt-strap{font-size:26px;margin-top:14px;padding-top:12px;max-width:640px}.bdt-hero .bdt-when{font-size:14px;margin-top:8px}.bdt-hero .bdt-when svg{width:16px;height:16px}}</style>"""
    return hd.replace('</head>', strap_css + '</head>', 1)


if __name__ == '__main__':
    for name, fn in [('bdt_nicknames.html', nicknames), ('bdt_tour_videos.html', videos), ('bdt_rules.html', rules), ('bdt_tour_2027.html', tour_2027)]:
        h = fn()
        open(os.path.join(OUT, name), 'w', encoding='utf-8').write(h)
        print('wrote', name, len(h) // 1024, 'KB')
